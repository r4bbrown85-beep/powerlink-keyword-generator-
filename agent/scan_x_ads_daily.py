# -*- coding: utf-8 -*-
"""
X(트위터) Ads API 계정별 일단위 집행금액(BILLING) 증분 스캔 + 구글시트 업로드.
매일 새벽 스케줄 작업으로 실행 (독립 프로세스 — 레이트리밋 걸리면 그냥 sleep하고 계속 진행,
Claude Code 세션 안에서 겪었던 백그라운드 강제종료 이슈와 무관).

전략: 이번 달 확정된 과거 데이터는 그대로 두고, 최근 며칠(RECHECK_DAYS)만 다시 조회해서
덮어쓴다 — 매일 전체 월을 처음부터 훑지 않아 실행시간이 날짜에 관계없이 일정하다
([[project_x_ads_api_integration]] 참고, 2026-09-10 팀장 지시로 전체재조회->증분 방식 전환).

547개 계정 전체는 매일 funding_instruments만 가볍게 재확인(신규 집행 계정 감지 목적).

구글시트 업로드 방식 (2026-09-15 전면 개편): 여러 명이 같이 편집하는 공유 시트라 "매일
클리어 후 총액순 재정렬해서 통째로 재작성"하는 방식을 버리고, 기존 행/순서/서식/수동 입력
(참조 광고주명, 계정명 뒤 대행사 표기, 색깔, 화이트리스트 O/X 수동 조정 등)은 절대 건드리지
않은 채 어제자 날짜 열만 계정ID 매칭으로 추가하는 증분 방식으로 전환. 고정 행 번호 대신
매 실행마다 시트를 읽어 KRW/USD 블록 경계("<USD 운영계정>" 마커)를 동적으로 찾는다 - 팀원이
수동으로 행을 넣거나 지워도 안전. 신규로 매출이 잡힌 계정은 각 블록 맨 밑에 추가하고, 버퍼가
모자라면 그 자리에 행을 삽입해서 아래 블록을 통째로 밀어낸다.
"""
import sys, os, json, time, re, calendar
from datetime import datetime, timedelta, date
from zoneinfo import ZoneInfo

sys.stdout.reconfigure(encoding='utf-8')

from dotenv import load_dotenv
from requests_oauthlib import OAuth1
import requests
import urllib3
urllib3.disable_warnings()

from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from openpyxl.utils import get_column_letter

BASE_DIR = r"C:\Users\Administrator\Desktop\powerlink_keyword_generator"
load_dotenv(os.path.join(BASE_DIR, ".env"))

AUTH = OAuth1(
    os.environ["X_API_KEY"],
    os.environ["X_API_KEY_SECRET"],
    os.environ["X_ACCESS_TOKEN"],
    os.environ["X_ACCESS_TOKEN_SECRET"],
)

ACCOUNTS_PATH = r"C:/Users/Administrator/AppData/Local/Temp/x_ads_accounts.json"
STORE_PATH = r"C:/Users/Administrator/Desktop/00. 클로드코드 생성물/네이버 GFA API/x_ads_daily_store.json"
EXCEL_OUT = r"C:/Users/Administrator/Desktop/00. 클로드코드 생성물/네이버 GFA API/X_이번달_계정별_일단위_매출.xlsx"

# 2026-10-01: X 트래킹 전용 스프레드시트(네이버 매출 현황과 같은 파일)로 이전, 매월 새 탭
# 구조로 전환 - 더 이상 "26년 주간회의록" 시트의 고정 탭 2개(원본/화이트리스트)를 쓰지 않는다.
# 화이트리스트 적용 탭 하나만 운영(2026-10-01 팀장 확정). 탭 이름은 x_tab_name()으로 매번 계산.
SHEET_ID = "1V-gAyGok-H29rzXtvpE5UTFjhh-DYNbL6SW74dBrwNs"
WHITELIST_PATH = r"C:/Users/Administrator/Desktop/00. 클로드코드 생성물/X API관련/x_invoiced_whitelist_2026_1to8.json"


def x_tab_name(d):
    return f"{d.month}월_X일단위 매출 트래킹(화이트리스트 적용)"
TOKEN_PATH = os.path.join(BASE_DIR, "config", "briefing_token.json")
# 스코프 하드코딩 금지 - briefing_token.json을 morning_briefing.py 등과 공유하므로
# get_creds() 호출부에서 scopes 인자를 생략해 토큰 파일의 실제 스코프를 그대로 쓴다
# (2026-09-11: 스코프 불일치로 아침 브리핑이 크래시난 사고 이후 전체 통일).

KRW_START = 7  # KRW 블록은 항상 7행부터 시작(고정). 끝나는 지점은 "<USD 운영계정>" 마커로 매번 동적 탐색.
USD_KRW_RATE = 1300
HEADER_ROW = 6

# 열 레이아웃은 더 이상 하드코딩하지 않는다(2026-09-23). 팀원이 시트 왼쪽에 새 열
# ("정산 금액(추정)" 등)을 끼워넣거나 메모를 추가해도 매 실행마다 헤더(6행) 텍스트로
# 실제 위치를 다시 찾는다 - detect_columns() 참고.


def load_whitelist_ids():
    with open(WHITELIST_PATH, encoding="utf-8") as f:
        wl = json.load(f)
    return set(wl["accounts"].keys())

RECHECK_DAYS = 3  # 최근 N일은 매번 재조회해서 덮어씀 (뒤늦게 반영되는 정산 보정 대비)
BASE_DELAY = 3.8


def log(msg):
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}", flush=True)


# ── X Ads API ────────────────────────────────────────────
def api_get(url, params, retry=5):
    for attempt in range(retry):
        try:
            resp = requests.get(url, auth=AUTH, params=params, verify=False, timeout=20)
        except Exception as e:
            log(f"  [네트워크 오류] {type(e).__name__}: {e} - 5초 후 재시도")
            time.sleep(5)
            continue

        remaining = resp.headers.get("x-rate-limit-remaining")
        reset = resp.headers.get("x-rate-limit-reset")
        if resp.status_code == 429 or (remaining is not None and int(remaining) <= 1):
            wait = max(int(reset) - int(time.time()), 5) if reset else 60
            log(f"  [레이트리밋] {wait + 2}초 대기")
            time.sleep(wait + 2)
            continue
        if resp.status_code == 200:
            time.sleep(BASE_DELAY)
            return resp.json()
        if resp.status_code >= 500:
            log(f"  [서버 오류 {resp.status_code}] 재시도")
            time.sleep(5)
            continue
        return {"_error": True, "status": resp.status_code, "body": resp.text[:300]}
    return {"_error": True, "status": -1, "body": "재시도 초과"}


def refresh_accounts_list():
    """X Ads API에서 접근 가능한 전체 계정 목록을 매번 다시 받아와 ACCOUNTS_PATH에 저장한다.
    2026-09-08에 한 번 긁은 스냅샷을 그 뒤로 한 번도 갱신 안 하고 계속 재사용해온 게
    발견되어(2026-10-01) 추가 - 9/8 이후 X 쪽에서 새로 Account Access를 열어준 계정은
    지금까지 전부 사각지대였다. 목록 조회는 funding_instruments 547콜에 비해 가벼우니
    매일 실행 때마다 갱신한다(월 1회가 아니라 매번 - 더 단순하고 더 안전)."""
    url = "https://ads-api.x.com/12/accounts"
    all_accounts = []
    cursor = None
    while True:
        params = {"cursor": cursor} if cursor else {}
        body = api_get(url, params)
        if body.get("_error"):
            raise RuntimeError(f"계정 목록 조회 실패: {body}")
        all_accounts.extend(body.get("data", []))
        cursor = body.get("next_cursor")
        if not cursor:
            break

    old_ids = set()
    if os.path.exists(ACCOUNTS_PATH):
        with open(ACCOUNTS_PATH, encoding="utf-8") as f:
            old_ids = {a["id"] for a in json.load(f)}
    new_ids = {a["id"] for a in all_accounts}
    added, removed = new_ids - old_ids, old_ids - new_ids

    with open(ACCOUNTS_PATH, "w", encoding="utf-8") as f:
        json.dump(all_accounts, f, ensure_ascii=False, indent=2)

    if added or removed:
        log(f"계정 목록 갱신: {len(all_accounts)}개 (신규 +{len(added)}개, 제외 -{len(removed)}개)")
        if added:
            added_names = {a["id"]: a["name"] for a in all_accounts if a["id"] in added}
            for aid in added:
                log(f"  [신규 접근권한] {aid} {added_names.get(aid, '')}")
    else:
        log(f"계정 목록 갱신: {len(all_accounts)}개 (변동 없음)")
    return all_accounts


def get_active_funding_instruments(account_id):
    url = f"https://ads-api.x.com/12/accounts/{account_id}/funding_instruments"
    body = api_get(url, {})
    if body.get("_error"):
        return None, body
    fis = [
        fi for fi in body.get("data", [])
        if fi.get("entity_status") == "ACTIVE" and not fi.get("deleted") and fi.get("able_to_fund")
    ]
    return fis, None


def fmt_utc(dt):
    return dt.astimezone(ZoneInfo("UTC")).strftime("%Y-%m-%dT%H:%M:%SZ")


def get_billing_days(account_id, fi_id, tz_name, start_date, end_date_exclusive):
    """[start_date, end_date_exclusive) 구간의 일별 billed_charge_local_micro 조회 (7일 이내 가정)"""
    tz = ZoneInfo(tz_name)
    start_local = datetime.combine(start_date, datetime.min.time(), tzinfo=tz)
    end_local = datetime.combine(end_date_exclusive, datetime.min.time(), tzinfo=tz)
    url = f"https://ads-api.x.com/12/stats/accounts/{account_id}"
    params = {
        "entity": "FUNDING_INSTRUMENT",
        "entity_ids": fi_id,
        "metric_groups": "BILLING",
        "granularity": "DAY",
        "placement": "ALL_ON_TWITTER",
        "start_time": fmt_utc(start_local),
        "end_time": fmt_utc(end_local),
    }
    body = api_get(url, params)
    daily = {}
    if body.get("_error"):
        log(f"    [실패] fi={fi_id}: {body}")
        return daily
    try:
        id_data = body["data"][0]["id_data"][0]
        metrics = id_data.get("metrics") or {}
        charges = metrics.get("billed_charge_local_micro") or []
    except (KeyError, IndexError):
        charges = []
    n_days = (end_date_exclusive - start_date).days
    for i in range(n_days):
        day = (start_date + timedelta(days=i)).isoformat()
        val = charges[i] if i < len(charges) and charges[i] is not None else 0
        daily[day] = val
    return daily


# ── 저장소 ───────────────────────────────────────────────
def load_store():
    if os.path.exists(STORE_PATH):
        with open(STORE_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {"month": None, "accounts": {}}


def save_store(store):
    with open(STORE_PATH, "w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False, indent=2)


def save_store_to(path, store):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False, indent=2)


# ── 메인 ─────────────────────────────────────────────────
def main():
    today = date.today()
    yesterday = today - timedelta(days=1)
    month_start = yesterday.replace(day=1)  # "어제" 기준 월 — 매월 1일 실행 시 전월 마지막날(어제)을 그 달 store로 정확히 반영하기 위함

    recheck_start = max(month_start, yesterday - timedelta(days=RECHECK_DAYS - 1))

    accounts = refresh_accounts_list()

    store = load_store()
    cur_month_key = month_start.isoformat()
    if store.get("month") != cur_month_key:
        if store.get("month"):
            archive_path = STORE_PATH.replace(".json", f"_{store['month']}_archive.json")
            save_store_to(archive_path, store)
            log(f"월 변경 감지 - 이전 달 데이터 아카이브: {archive_path}")
        store = {"month": cur_month_key, "accounts": {}}

    log(f"시작 - 조회범위: {recheck_start} ~ {yesterday} (최근 {RECHECK_DAYS}일 재확인), 전체 계정 {len(accounts)}개")

    for i, acc in enumerate(accounts, start=1):
        acc_id = acc["id"]
        fis, err = get_active_funding_instruments(acc_id)
        if err:
            store["accounts"][acc_id] = {"name": acc["name"], "error": str(err)}
            continue
        if not fis:
            store["accounts"].setdefault(acc_id, {"name": acc["name"], "active_funding_instruments": 0, "by_currency": {}})
            store["accounts"][acc_id]["active_funding_instruments"] = 0
            continue

        entry = store["accounts"].setdefault(acc_id, {"name": acc["name"], "by_currency": {}})
        entry["name"] = acc["name"]
        entry["active_funding_instruments"] = len(fis)
        by_currency = entry.setdefault("by_currency", {})

        for fi in fis:
            cur = fi.get("currency") or "UNKNOWN"
            fresh = get_billing_days(acc_id, fi["id"], acc.get("timezone") or "UTC", recheck_start, yesterday + timedelta(days=1))
            bucket = by_currency.setdefault(cur, {})
            for d, v in fresh.items():
                bucket[d] = v  # 최근 N일은 덮어쓰기 (증분)

        if i % 20 == 0 or i == len(accounts):
            save_store(store)
            log(f"진행 {i}/{len(accounts)}")

    save_store(store)
    log("스캔 완료. 엑셀/시트 갱신 시작")

    all_dates = date_range(month_start, yesterday)
    write_excel(build_rows(store, all_dates), month_start)

    svc = get_sheets_service()
    whitelist_ids = load_whitelist_ids()

    sheet_tab, _ = ensure_month_tab(svc, month_start)
    cols = detect_columns(svc, sheet_tab)
    recheck_dates = []
    d = recheck_start
    while d <= yesterday:
        recheck_dates.append(d)
        d += timedelta(days=1)
    update_fixed_grid_tab(svc, sheet_tab, store, all_dates, recheck_dates, cols, whitelist_ids=whitelist_ids)

    log("전체 완료")


def date_range(month_start, yesterday):
    dates = []
    d = month_start
    while d <= yesterday:
        dates.append(d.isoformat())
        d += timedelta(days=1)
    return dates


def build_rows(store, dates):
    """로컬 엑셀 산출물 전용 - 총액순 정렬된 전체 계정 리스트(구글시트에는 안 씀)."""
    rows = []
    for acc_id, v in store["accounts"].items():
        for cur, days in (v.get("by_currency") or {}).items():
            total = sum(days.get(d, 0) for d in dates) / 1_000_000
            if total <= 0:
                continue
            daily_vals = [(days.get(d, 0) / 1_000_000) if days.get(d) else None for d in dates]
            rows.append({
                "account_id": acc_id, "name": v["name"], "currency": cur,
                "total": total, "avg": total / len(dates), "daily": daily_vals,
            })
    rows.sort(key=lambda r: -r["total"])
    return rows, dates


def write_excel(rows_dates, month_start):
    rows, dates = rows_dates
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "X 계정별 일단위"
    ws.cell(2, 2, f"● {month_start.month}월 X(트위터) 계정별 일단위 집행금액 트래킹")
    ws.cell(2, 2).font = Font(bold=True, size=13)

    header_row = 4
    headers = ["광고계정ID", "계정명", "통화", "총계", "일 평균 금액"] + [d[5:].replace("-", "/") for d in dates]
    for j, h in enumerate(headers, start=2):
        c = ws.cell(header_row, j, h)
        c.fill = PatternFill("solid", fgColor="4F81BD")
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rows):
        row = header_row + 1 + i
        ws.cell(row, 2, r["account_id"])
        ws.cell(row, 3, r["name"])
        ws.cell(row, 4, r["currency"])
        ws.cell(row, 5, round(r["total"], 2))
        ws.cell(row, 6, round(r["avg"], 2))
        for k, val in enumerate(r["daily"]):
            ws.cell(row, 7 + k, round(val, 2) if val else None)

    widths = {2: 14, 3: 34, 4: 8, 5: 14, 6: 14}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for k in range(len(dates)):
        ws.column_dimensions[get_column_letter(7 + k)].width = 12

    wb.save(EXCEL_OUT)
    log(f"엑셀 저장: {EXCEL_OUT} ({len(rows)}행)")


def get_sheets_service():
    creds = Credentials.from_authorized_user_file(TOKEN_PATH)  # scopes 생략 = 토큰 파일에 저장된 스코프 그대로 사용
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        with open(TOKEN_PATH, "w", encoding="utf-8") as f:
            f.write(creds.to_json())
    return build('sheets', 'v4', credentials=creds)


def get_sheet_id(svc, sheet_tab):
    meta = svc.spreadsheets().get(spreadsheetId=SHEET_ID, fields="sheets(properties(sheetId,title))").execute()
    return next(s["properties"]["sheetId"] for s in meta["sheets"] if s["properties"]["title"] == sheet_tab)


def _norm(s):
    return (s or "").replace("\n", "").replace(" ", "").strip()


def detect_columns(svc, sheet_tab, max_col=60):
    """헤더(6행) 텍스트로 계정ID/계정명/통화/총계/일평균/화이트리스트/첫날짜 열 위치를 매번
    새로 찾는다. 팀원이 왼쪽에 새 열(예: '정산 금액(추정)')을 끼워넣어 기존 열이 밀려도
    하드코딩된 열 번호에 의존하지 않으므로 안전하다(2026-09-23, 실제 이 문제로 화이트리스트
    탭 레이아웃이 2칸 밀린 채 발견되어 도입)."""
    last_col = get_column_letter(max_col)
    resp = svc.spreadsheets().values().get(
        spreadsheetId=SHEET_ID, range=f"'{sheet_tab}'!A{HEADER_ROW}:{last_col}{HEADER_ROW}"
    ).execute()
    row = resp.get("values", [[]])
    row = row[0] if row else []

    id_col = name_col = curr_col = total_col = avg_col = flag_col = None
    for idx0, val in enumerate(row):
        v = _norm(val)
        if v == "광고계정ID":
            id_col = idx0 + 1
        elif v == "계정명":
            name_col = idx0 + 1
        elif v == "통화":
            curr_col = idx0 + 1
        elif v == "총계":
            total_col = idx0 + 1
        elif "일평균" in v:
            avg_col = idx0 + 1
        elif "화이트리스트" in v:
            flag_col = idx0 + 1

    required = {"광고계정ID": id_col, "계정명": name_col, "통화": curr_col, "총계": total_col, "일평균금액": avg_col}
    missing = [name for name, col in required.items() if col is None]
    if missing:
        raise RuntimeError(f"{sheet_tab}: 헤더({HEADER_ROW}행)에서 {missing} 열을 못 찾음 - 시트 구조가 크게 바뀐 것으로 보임, 확인 필요")

    # 첫 날짜 열: 보통 일평균 바로 다음 열이지만, 혹시 그 사이에 다른 열이 끼어들 경우를
    # 대비해 "MM/DD" 형태 헤더가 실제로 처음 나오는 열을 우선 탐색한다.
    date_re = re.compile(r"^\d{2}/\d{2}$")
    first_date_col = avg_col + 1
    for idx0, val in enumerate(row):
        col1 = idx0 + 1
        if col1 <= avg_col:
            continue
        if date_re.match((val or "").strip()):
            first_date_col = col1
            break

    cols = dict(id=id_col, name=name_col, curr=curr_col, total=total_col, avg=avg_col,
                first_date=first_date_col, flag=flag_col)
    log(f"{sheet_tab}: 열 위치 탐지 - {cols}")
    return cols


# ── 월별 탭 관리(2026-10-01 신규) ─────────────────────────
# 날짜 열을 매일 하나씩 "추가"하던 예전 방식과 달리, 이제는 탭 하나에 그 달 전체
# 날짜(1일~말일) 열이 처음부터 다 만들어져 있고 수식(F열 총계, 5행/USD서브토탈행 합계)도
# 전부 미리 들어있다 - 팀장이 직접 만든 10월 탭 실측 확인. 그래서 스크립트는 더 이상 열을
# 추가하지 않고, 정해진 날짜의 열 위치를 "그 달 n일째 = 첫날짜열 + (n-1)"로 바로 계산해서
# 그 자리에 값만 쓴다.

def ensure_month_tab(svc, month_start):
    """해당 월의 탭이 있으면 그대로 쓰고, 없으면 바로 전월 탭을 복제해서 새로 만든다
    (전월 탭의 서식·수식·상단 수동 박스 구조를 그대로 물려받음). 새 탭은 전월 탭
    바로 왼쪽에 배치한다(2026-10-01 팀장 지시)."""
    tab_name = x_tab_name(month_start)
    meta = svc.spreadsheets().get(spreadsheetId=SHEET_ID, fields="sheets(properties(sheetId,title,index))").execute()
    sheets = meta["sheets"]
    existing = next((s for s in sheets if s["properties"]["title"] == tab_name), None)
    if existing:
        return tab_name, existing["properties"]["sheetId"]

    prev_month = (month_start - timedelta(days=1)).replace(day=1)
    prev_tab_name = x_tab_name(prev_month)
    prev_sheet = next((s for s in sheets if s["properties"]["title"] == prev_tab_name), None)
    if not prev_sheet:
        raise RuntimeError(f"{tab_name} 탭이 없고, 복제할 전월 탭({prev_tab_name})도 못 찾음 - 수동으로 만들어야 함")

    prev_sheet_id = prev_sheet["properties"]["sheetId"]
    prev_index = prev_sheet["properties"]["index"]

    dup_resp = svc.spreadsheets().batchUpdate(spreadsheetId=SHEET_ID, body={"requests": [
        {"duplicateSheet": {"sourceSheetId": prev_sheet_id, "insertSheetIndex": prev_index, "newSheetName": tab_name}}
    ]}).execute()
    new_sheet_id = dup_resp["replies"][0]["duplicateSheet"]["properties"]["sheetId"]
    log(f"신규 탭 생성: '{tab_name}' ('{prev_tab_name}' 복제, 그 왼쪽에 배치)")

    _reset_month_tab_contents(svc, tab_name, month_start)
    log(f"'{tab_name}': 계정 데이터·날짜 값 초기화 + 날짜 헤더를 {month_start.month}월 기준으로 재작성 완료 "
        f"- 상단 1~3행 수동 박스(예약형 Takeover 등)는 전월 내용 그대로이니 월별 입력값 검토 필요")
    return tab_name, new_sheet_id


def _reset_month_tab_contents(svc, tab_name, month_start):
    """복제 직후의 새 탭에서: 계정 데이터(ID/계정명/통화/화이트리스트)와 날짜별 값을 전부
    지우고(=그 달 다시 새로 시작), 날짜 헤더(6행)와 제목(4행)만 새 달 기준으로 다시 쓴다.
    F열(총계)·5행/USD서브토탈행 합계 수식은 복제본에 이미 있으므로 건드리지 않는다."""
    cols = detect_columns(svc, tab_name)
    layout = find_tab_layout(svc, tab_name, cols)
    id_col = get_column_letter(cols["id"])
    flag_col = get_column_letter(cols["flag"]) if cols.get("flag") else get_column_letter(cols["curr"])
    first_date_col = get_column_letter(cols["first_date"])
    last_date_col_idx1 = cols["first_date"] + 30  # 31일치까지 항상 확보(어느 달을 복제해왔든 다음 31일짜리 달도 안전하게)
    last_date_col = get_column_letter(last_date_col_idx1)

    # 복제해온 탭의 그리드가 31일치보다 좁으면(예: 28~30일짜리 달을 복제한 경우) 먼저 넓힌다
    meta = svc.spreadsheets().get(spreadsheetId=SHEET_ID, fields="sheets(properties(sheetId,title,gridProperties))").execute()
    sheet_meta = next(s for s in meta["sheets"] if s["properties"]["title"] == tab_name)
    current_cols = sheet_meta["properties"]["gridProperties"].get("columnCount", 26)
    if last_date_col_idx1 > current_cols:
        add = last_date_col_idx1 - current_cols
        svc.spreadsheets().batchUpdate(spreadsheetId=SHEET_ID, body={"requests": [{
            "appendDimension": {"sheetId": sheet_meta["properties"]["sheetId"], "dimension": "COLUMNS", "length": add}
        }]}).execute()

    clear_ranges = [
        f"'{tab_name}'!{id_col}{KRW_START}:{flag_col}{layout['last_krw_row']}",
        f"'{tab_name}'!{id_col}{layout['usd_start_row']}:{flag_col}{layout['last_usd_row']}",
        f"'{tab_name}'!{first_date_col}{KRW_START}:{last_date_col}{layout['last_krw_row']}",
        f"'{tab_name}'!{first_date_col}{layout['usd_start_row']}:{last_date_col}{layout['last_usd_row']}",
    ]
    svc.spreadsheets().values().batchClear(spreadsheetId=SHEET_ID, body={"ranges": clear_ranges}).execute()

    days_in_month = calendar.monthrange(month_start.year, month_start.month)[1]
    header_data = [{"range": f"'{tab_name}'!B4", "values": [[f"● {month_start.month}월 X(트위터) 계정별 일단위 집행금액 트래킹"]]}]
    for dd in range(1, 32):
        col = get_column_letter(cols["first_date"] + dd - 1)
        label = f"{month_start.month:02d}/{dd:02d}" if dd <= days_in_month else ""
        header_data.append({"range": f"'{tab_name}'!{col}6", "values": [[label]]})
    svc.spreadsheets().values().batchUpdate(
        spreadsheetId=SHEET_ID, body={"valueInputOption": "RAW", "data": header_data}
    ).execute()


def find_tab_layout(svc, sheet_tab, cols):
    """클리어를 하지 않는 증분 방식이므로, 매 실행마다 현재 시트를 읽어 KRW/USD 블록의
    실제 계정ID->행 매핑과 경계("<USD 운영계정>" 마커)를 동적으로 찾는다. 고정 행 상수에
    의존하지 않아 팀원이 행을 수동으로 넣거나 지워도 안전하다."""
    id_col = get_column_letter(cols["id"])
    curr_col = get_column_letter(cols["curr"])
    name_rel = cols["name"] - cols["id"]
    curr_rel = cols["curr"] - cols["id"]

    resp = svc.spreadsheets().values().get(
        spreadsheetId=SHEET_ID, range=f"'{sheet_tab}'!{id_col}{KRW_START}:{curr_col}400"
    ).execute()
    rows = resp.get("values", [])

    krw_id_to_row = {}
    last_krw_row = KRW_START - 1
    usd_subtotal_row = None

    for i, row in enumerate(rows):
        row_num = KRW_START + i
        cell_id = row[0] if len(row) > 0 else ""
        cell_name = row[name_rel] if len(row) > name_rel else ""
        cell_curr = row[curr_rel] if len(row) > curr_rel else ""
        if cell_id == "<USD 운영계정>":
            usd_subtotal_row = row_num
            break
        if cell_id and cell_name and cell_curr:
            krw_id_to_row[cell_id] = row_num
            last_krw_row = row_num
        # 빈 버퍼 행이나 불완전한 행(과거 메모 등)은 건너뜀 - last_krw_row는 갱신 안 함

    if usd_subtotal_row is None:
        raise RuntimeError(f"{sheet_tab}: '<USD 운영계정>' 마커를 못 찾음 - 시트 구조 확인 필요")

    usd_total_row = usd_subtotal_row + 1
    usd_start_row = usd_subtotal_row + 2

    resp2 = svc.spreadsheets().values().get(
        spreadsheetId=SHEET_ID, range=f"'{sheet_tab}'!{id_col}{usd_start_row}:{curr_col}{usd_start_row + 100}"
    ).execute()
    rows2 = resp2.get("values", [])
    usd_id_to_row = {}
    last_usd_row = usd_start_row - 1
    for i, row in enumerate(rows2):
        row_num = usd_start_row + i
        cell_id = row[0] if len(row) > 0 else ""
        cell_name = row[name_rel] if len(row) > name_rel else ""
        cell_curr = row[curr_rel] if len(row) > curr_rel else ""
        if cell_id and cell_name and cell_curr:
            usd_id_to_row[cell_id] = row_num
            last_usd_row = row_num
        elif not cell_id and not cell_name and not cell_curr:
            break  # USD 블록 끝(공백행) - 그 아래엔 아무것도 없다고 가정

    return {
        "krw_id_to_row": krw_id_to_row, "last_krw_row": last_krw_row,
        "usd_id_to_row": usd_id_to_row, "last_usd_row": last_usd_row,
        "usd_subtotal_row": usd_subtotal_row, "usd_total_row": usd_total_row, "usd_start_row": usd_start_row,
    }


def insert_rows(svc, sheet_tab, before_row1, count):
    """before_row1(1-based) 앞에 count개의 빈 행을 삽입 - 그 아래 내용(수식·서식 포함) 전부
    자동으로 밀린다(구글시트 기본 동작)."""
    meta = svc.spreadsheets().get(spreadsheetId=SHEET_ID, fields="sheets(properties(sheetId,title))").execute()
    sheet_id = next(s["properties"]["sheetId"] for s in meta["sheets"] if s["properties"]["title"] == sheet_tab)
    svc.spreadsheets().batchUpdate(spreadsheetId=SHEET_ID, body={"requests": [{
        "insertDimension": {
            "range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": before_row1 - 1, "endIndex": before_row1 - 1 + count},
            "inheritFromBefore": False,
        }
    }]}).execute()


def get_nonkrw_currency(store, acc_id):
    days_by_cur = (store["accounts"].get(acc_id, {}).get("by_currency")) or {}
    for cur in days_by_cur:
        if cur != "KRW":
            return cur
    return "USD"


def account_value(store, acc_id, currency, date_str):
    acc = store["accounts"].get(acc_id)
    if not acc:
        return None
    days = (acc.get("by_currency") or {}).get(currency)
    if not days:
        return None
    v = days.get(date_str)
    if not v:
        return None
    return v / 1_000_000


def account_cum_total(store, acc_id, currency, dates):
    acc = store["accounts"].get(acc_id)
    if not acc:
        return 0
    days = (acc.get("by_currency") or {}).get(currency) or {}
    return sum(days.get(d, 0) for d in dates) / 1_000_000


def apply_new_column_formatting(svc, sheet_tab, first_new_col_idx1, last_new_col_idx1, krw_start, krw_end, usd_start, usd_end):
    meta = svc.spreadsheets().get(spreadsheetId=SHEET_ID, fields="sheets(properties(sheetId,title))").execute()
    sheet_id = next(s["properties"]["sheetId"] for s in meta["sheets"] if s["properties"]["title"] == sheet_tab)
    requests_ = [
        {"repeatCell": {
            "range": {"sheetId": sheet_id, "startRowIndex": 5, "endRowIndex": 6,
                       "startColumnIndex": first_new_col_idx1 - 1, "endColumnIndex": last_new_col_idx1},
            "cell": {"userEnteredFormat": {
                "backgroundColor": {"red": 0.30980393, "green": 0.5058824, "blue": 0.7411765},
                "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                "horizontalAlignment": "CENTER", "wrapStrategy": "WRAP",
            }},
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,wrapStrategy)",
        }},
        {"repeatCell": {
            "range": {"sheetId": sheet_id, "startRowIndex": krw_start - 1, "endRowIndex": krw_end,
                       "startColumnIndex": first_new_col_idx1 - 1, "endColumnIndex": last_new_col_idx1},
            "cell": {"userEnteredFormat": {"numberFormat": {"type": "NUMBER", "pattern": "#,##0"}}},
            "fields": "userEnteredFormat.numberFormat",
        }},
        {"repeatCell": {
            "range": {"sheetId": sheet_id, "startRowIndex": usd_start - 1, "endRowIndex": usd_end,
                       "startColumnIndex": first_new_col_idx1 - 1, "endColumnIndex": last_new_col_idx1},
            "cell": {"userEnteredFormat": {"numberFormat": {"type": "NUMBER", "pattern": "#,##0"}}},
            "fields": "userEnteredFormat.numberFormat",
        }},
    ]
    svc.spreadsheets().batchUpdate(spreadsheetId=SHEET_ID, body={"requests": requests_}).execute()


def apply_whitelist_conditional_format(svc, sheet_tab, flag_col_idx1, krw_start, krw_end, usd_start, usd_end):
    """화이트리스트 열 O=초록/X=빨강 조건부 서식 - 매번 규칙을 지우고 다시 걸어서
    규칙이 중복 누적되지 않게 한다."""
    meta = svc.spreadsheets().get(spreadsheetId=SHEET_ID, fields="sheets(properties(sheetId,title),conditionalFormats)").execute()
    sheet = next(s for s in meta["sheets"] if s["properties"]["title"] == sheet_tab)
    sheet_id = sheet["properties"]["sheetId"]
    n_existing_rules = len(sheet.get("conditionalFormats", []))

    def rng(r1, r2, c1, c2):
        return {"sheetId": sheet_id, "startRowIndex": r1 - 1, "endRowIndex": r2, "startColumnIndex": c1 - 1, "endColumnIndex": c2}

    if krw_end < krw_start or usd_end < usd_start:
        return

    requests_ = [{"deleteConditionalFormatRule": {"sheetId": sheet_id, "index": i}} for i in reversed(range(n_existing_rules))]
    requests_ += [
        {"addConditionalFormatRule": {"rule": {
            "ranges": [rng(krw_start, krw_end, flag_col_idx1, flag_col_idx1), rng(usd_start, usd_end, flag_col_idx1, flag_col_idx1)],
            "booleanRule": {"condition": {"type": "TEXT_EQ", "values": [{"userEnteredValue": "X"}]},
                             "format": {"backgroundColor": {"red": 0.9764706, "green": 0.84705883, "blue": 0.84705883}}}
        }, "index": 0}},
        {"addConditionalFormatRule": {"rule": {
            "ranges": [rng(krw_start, krw_end, flag_col_idx1, flag_col_idx1), rng(usd_start, usd_end, flag_col_idx1, flag_col_idx1)],
            "booleanRule": {"condition": {"type": "TEXT_EQ", "values": [{"userEnteredValue": "O"}]},
                             "format": {"backgroundColor": {"red": 0.84705883, "green": 0.9490196, "blue": 0.84705883}}}
        }, "index": 0}},
    ]
    svc.spreadsheets().batchUpdate(spreadsheetId=SHEET_ID, body={"requests": requests_}).execute()


def update_fixed_grid_tab(svc, sheet_tab, store, all_dates, recheck_dates, cols, whitelist_ids=None):
    """고정 그리드 탭(그 달 1~말일 날짜 열이 이미 다 만들어져 있고 F열 총계·5행/USD서브토탈행
    합계 수식도 이미 들어있는 구조, 2026-10-01 신규)에 값만 쓴다. 날짜 열을 추가하지 않고
    "그 달 n일째 = 첫날짜열 + (n-1)"로 열 위치를 바로 계산한다. 기존 행 순서·색·수식은 그대로
    두고, 신규로 매출이 잡힌 계정만 KRW/USD 각 블록 맨 밑에 추가한다(모자라면 행 삽입)."""
    if not recheck_dates:
        log(f"{sheet_tab}: 갱신할 날짜 없음")
        return

    id_col = get_column_letter(cols["id"])
    name_col = get_column_letter(cols["name"])
    curr_col = get_column_letter(cols["curr"])
    total_col = get_column_letter(cols["total"])
    avg_col = get_column_letter(cols["avg"])
    first_date_col_idx1 = cols["first_date"]
    flag_col = get_column_letter(cols["flag"]) if cols.get("flag") else None

    layout = find_tab_layout(svc, sheet_tab, cols)
    krw_id_to_row = layout["krw_id_to_row"]
    usd_id_to_row = layout["usd_id_to_row"]
    last_krw_row = layout["last_krw_row"]
    last_usd_row = layout["last_usd_row"]
    usd_subtotal_row = layout["usd_subtotal_row"]
    usd_start_row = layout["usd_start_row"]

    tracked_krw_ids = set(krw_id_to_row)
    tracked_usd_ids = set(usd_id_to_row)

    new_krw, new_usd = [], []
    for acc_id, v in store["accounts"].items():
        for cur, days in (v.get("by_currency") or {}).items():
            total_cum = sum(days.get(d, 0) for d in all_dates)
            if total_cum <= 0:
                continue
            if cur == "KRW":
                if acc_id not in tracked_krw_ids:
                    new_krw.append((acc_id, v["name"]))
            else:
                if acc_id not in tracked_usd_ids:
                    new_usd.append((acc_id, v["name"]))

    month_ref = date.fromisoformat(all_dates[0]) if all_dates else date.today()
    days_in_month = calendar.monthrange(month_ref.year, month_ref.month)[1]
    last_date_col_idx1 = first_date_col_idx1 + days_in_month - 1

    # ---- KRW 신규 계정: 맨 밑에 추가, 버퍼 모자라면 그 자리에 행 삽입 ----
    if new_krw:
        available = usd_subtotal_row - (last_krw_row + 1)
        need_insert = max(0, len(new_krw) - available)
        if need_insert > 0:
            insert_row_at = last_krw_row + 1 + available  # 현재 "<USD 운영계정>" 행 위치
            insert_rows(svc, sheet_tab, insert_row_at, need_insert)
            usd_subtotal_row += need_insert
            usd_start_row += need_insert
            usd_id_to_row = {k: r + need_insert for k, r in usd_id_to_row.items()}
            last_usd_row += need_insert
            log(f"{sheet_tab}: KRW 블록 버퍼 부족 - {need_insert}행 삽입(USD 블록 자동으로 아래로 밀림)")
            # 삽입된 행은 템플릿 수식(F열 SUM)이 없으므로 직접 생성
            insert_formula_data = [
                {"range": f"'{sheet_tab}'!{total_col}{r}",
                 "values": [[f"=SUM({get_column_letter(first_date_col_idx1)}{r}:{get_column_letter(last_date_col_idx1)}{r})"]]}
                for r in range(insert_row_at, insert_row_at + need_insert)
            ]
            svc.spreadsheets().values().batchUpdate(
                spreadsheetId=SHEET_ID, body={"valueInputOption": "USER_ENTERED", "data": insert_formula_data}
            ).execute()
            # 삽입된 행은 inheritFromBefore=False라 숫자 서식도 비어있음 - 날짜 전 구간에 천단위 서식 적용
            apply_new_column_formatting(
                svc, sheet_tab, first_date_col_idx1, last_date_col_idx1,
                insert_row_at, insert_row_at + need_insert - 1, insert_row_at, insert_row_at - 1,
            )

        data = []
        for i, (acc_id, name) in enumerate(new_krw):
            row = last_krw_row + 1 + i
            data.append({"range": f"'{sheet_tab}'!{id_col}{row}", "values": [[acc_id]]})
            data.append({"range": f"'{sheet_tab}'!{name_col}{row}", "values": [[name]]})
            data.append({"range": f"'{sheet_tab}'!{curr_col}{row}", "values": [["KRW"]]})
            if flag_col:
                flag = "O" if whitelist_ids and acc_id in whitelist_ids else "X"
                data.append({"range": f"'{sheet_tab}'!{flag_col}{row}", "values": [[flag]]})
            krw_id_to_row[acc_id] = row
        svc.spreadsheets().values().batchUpdate(
            spreadsheetId=SHEET_ID, body={"valueInputOption": "USER_ENTERED", "data": data}
        ).execute()
        last_krw_row += len(new_krw)
        log(f"{sheet_tab}: 신규 KRW 계정 {len(new_krw)}개 추가")

    # ---- USD 신규 계정: 맨 밑에 추가 (USD 블록은 그 아래 아무것도 없어 버퍼 개념 없이 자유 성장) ----
    if new_usd:
        data = []
        for i, (acc_id, name) in enumerate(new_usd):
            row = last_usd_row + 1 + i
            cur = get_nonkrw_currency(store, acc_id)
            data.append({"range": f"'{sheet_tab}'!{id_col}{row}", "values": [[acc_id]]})
            data.append({"range": f"'{sheet_tab}'!{name_col}{row}", "values": [[name]]})
            data.append({"range": f"'{sheet_tab}'!{curr_col}{row}", "values": [[cur]]})
            data.append({"range": f"'{sheet_tab}'!{total_col}{row}",
                         "values": [[f"=SUM({get_column_letter(first_date_col_idx1)}{row}:{get_column_letter(last_date_col_idx1)}{row})"]]})
            if flag_col:
                flag = "O" if whitelist_ids and acc_id in whitelist_ids else "X"
                data.append({"range": f"'{sheet_tab}'!{flag_col}{row}", "values": [[flag]]})
            usd_id_to_row[acc_id] = row
        svc.spreadsheets().values().batchUpdate(
            spreadsheetId=SHEET_ID, body={"valueInputOption": "USER_ENTERED", "data": data}
        ).execute()
        last_usd_row += len(new_usd)
        log(f"{sheet_tab}: 신규 USD 계정 {len(new_usd)}개 추가")

    # ---- 날짜별 값 쓰기 - 열은 이미 다 있으므로(그 달 n일째 = 첫날짜열+(n-1)) 값만 덮어씀.
    # 헤더·합계 수식은 탭 생성 시점에 이미 들어있어 여기서 다시 안 건드린다.
    data_raw = []
    for d in recheck_dates:
        col = get_column_letter(first_date_col_idx1 + d.day - 1)
        d_str = d.isoformat()
        for acc_id, row in krw_id_to_row.items():
            val = account_value(store, acc_id, "KRW", d_str)
            data_raw.append({"range": f"'{sheet_tab}'!{col}{row}", "values": [[val if val else ""]]})
        for acc_id, row in usd_id_to_row.items():
            cur = get_nonkrw_currency(store, acc_id)
            val = account_value(store, acc_id, cur, d_str)
            data_raw.append({"range": f"'{sheet_tab}'!{col}{row}", "values": [[val if val else ""]]})
    if data_raw:
        svc.spreadsheets().values().batchUpdate(
            spreadsheetId=SHEET_ID, body={"valueInputOption": "RAW", "data": data_raw}
        ).execute()

    # ---- 일평균(G열) 전체 갱신 - 총계/날짜수 기반이라 매번 새로 계산해서 씀 ----
    n_dates_now = len(all_dates)
    avg_data = []
    for acc_id, row in krw_id_to_row.items():
        total_cum = account_cum_total(store, acc_id, "KRW", all_dates)
        avg_data.append({"range": f"'{sheet_tab}'!{avg_col}{row}", "values": [[round(total_cum / n_dates_now, 2)]]})
    for acc_id, row in usd_id_to_row.items():
        cur = get_nonkrw_currency(store, acc_id)
        total_cum = account_cum_total(store, acc_id, cur, all_dates)
        avg_data.append({"range": f"'{sheet_tab}'!{avg_col}{row}", "values": [[round(total_cum / n_dates_now, 2)]]})
    if avg_data:
        svc.spreadsheets().values().batchUpdate(
            spreadsheetId=SHEET_ID, body={"valueInputOption": "RAW", "data": avg_data}
        ).execute()

    if flag_col and (new_krw or new_usd):
        apply_whitelist_conditional_format(svc, sheet_tab, cols["flag"], KRW_START, last_krw_row, usd_start_row, last_usd_row)

    log(f"{sheet_tab}: {len(recheck_dates)}일치 값 갱신 완료 ({recheck_dates[0].isoformat()}~{recheck_dates[-1].isoformat()}), "
        f"신규계정 KRW {len(new_krw)}개/USD {len(new_usd)}개")


if __name__ == "__main__":
    main()
