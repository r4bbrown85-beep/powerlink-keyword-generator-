# -*- coding: utf-8 -*-
r"""네이버 일단위 매출 트래킹 — 4213_실적상세 CSV 처리 공용 모듈.

매일 팀장이 네이버 비즈센터에서 `4213_실적상세_YYYY-MM-DD_YYYY-MM-DD.csv`를 내려받아
경로만 주면 (1) 월별 매출현황 시트의 GFA/SA/테무GFA/테무SA 4개 컬럼을 채우고
(2) 일단위 매출 요약 xlsx를 만든다.

컬럼 정의는 2026-10-06에 10/1 발표치와 대조해 확정했다(전부 1원 단위 일치):
  GFA      = 성과형 DA 유상 매출 + 성과형 DA 유상 매출 보정 + ADVoost 쇼핑 유상 매출
  파워링크  = 웹방문(내부/PC, 내부/모바일, 외부/PC, 외부/모바일) + NCC유상매출/실적보정
  쇼핑검색  = 쇼핑(내부/PC, 내부/모바일, 외부/PC, 외부/모바일) + 쇼핑유상매출/실적보정
  SA       = 파워링크 + 쇼핑검색
  테무     = 광고주명이 'Elementary Innovation Pte. Ltd' (뒤 마침표 유무 2종 혼재 → 정규화 필수)

CSV 인코딩은 **cp949**이고 컬럼명 앞에 공백이 붙은 것이 많아 strip()이 필수다.

사용:
    from agent.naver_daily_report import compute_from_csv, update_sheet, read_sheet_series
    facts = compute_from_csv(csv_path)
    update_sheet(facts["daily"])          # dry_run=True로 먼저 확인 권장
리포트 본문 서술은 매일 달라지므로 build_excel(facts, narrative, out)에 문단을 넘겨 만든다.
"""
from __future__ import annotations

import json
import os
import re
from datetime import date, timedelta

import pandas as pd

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOKEN_PATH = os.path.join(BASE_DIR, "config", "briefing_token.json")
SHEET_ID = "1V-gAyGok-H29rzXtvpE5UTFjhh-DYNbL6SW74dBrwNs"
OUT_DIR = r"C:\Users\Administrator\Desktop\00. 클로드코드 생성물\네이버 일단위 매출 트래킹"
EPOCH = date(1899, 12, 30)
WD = ["월", "화", "수", "목", "금", "토", "일"]

GFA_COLS = ["성과형 DA 유상 매출", "성과형 DA 유상 매출 보정", "ADVoost 쇼핑 유상 매출"]
PL_COLS = ["웹방문(내부/PC)", "웹방문(내부/모바일)", "웹방문(외부/PC)",
           "웹방문(외부/모바일)", "NCC유상매출/실적보정"]
SH_COLS = ["쇼핑(내부/PC)", "쇼핑(내부/모바일)", "쇼핑(외부/PC)",
           "쇼핑(외부/모바일)", "쇼핑유상매출/실적보정"]
TEMU = "elementary innovation pte. ltd"


# ────────────────────────── CSV 집계 ──────────────────────────
def _norm_adv(s: str) -> str:
    """광고주명 정규화 — 같은 광고주가 뒤 마침표 유무로 두 줄로 갈리는 사례가 있다."""
    return re.sub(r"[.\s]+$", "", str(s).strip()).lower()


def compute_from_csv(csv_path: str) -> dict:
    df = pd.read_csv(csv_path, encoding="cp949", low_memory=False)
    df.columns = [c.strip() for c in df.columns]
    for c in GFA_COLS + PL_COLS + SH_COLS:
        if c not in df.columns:
            raise KeyError(f"컬럼 없음: {c}")
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

    df["_adv"] = df["광고주명"].map(_norm_adv)
    df["GFA"] = df[GFA_COLS].sum(axis=1)
    df["파워링크"] = df[PL_COLS].sum(axis=1)
    df["쇼핑검색"] = df[SH_COLS].sum(axis=1)
    df["SA"] = df["파워링크"] + df["쇼핑검색"]

    temu = df["_adv"] == TEMU
    df["테무GFA"] = df["GFA"].where(temu, 0)
    df["테무SA"] = df["SA"].where(temu, 0)

    daily = (df.groupby("날짜")[["GFA", "파워링크", "쇼핑검색", "SA", "테무GFA", "테무SA"]]
             .sum().round(0).astype("int64").sort_index())

    adv = (df.groupby(["날짜", "광고주명"])[["GFA", "SA"]].sum()
           .reset_index().rename(columns={"광고주명": "adv"}))

    return {
        "csv": csv_path,
        "daily": {d: daily.loc[d].to_dict() for d in daily.index},
        "adv": adv,
        "raw": df,
    }


def day_over_day(facts: dict, top: int = 5) -> dict:
    """마지막 날 vs 직전 날 광고주별 증감 + 신규 집행 광고주."""
    dates = sorted(facts["daily"])
    if len(dates) < 2:
        return {}
    last, prev = dates[-1], dates[-2]
    adv = facts["adv"]
    a = adv[adv["날짜"] == last].set_index("adv")[["GFA", "SA"]]
    b = adv[adv["날짜"] == prev].set_index("adv")[["GFA", "SA"]]
    j = a.join(b, how="outer", lsuffix="_last", rsuffix="_prev").fillna(0)
    out = {"last": last, "prev": prev}
    for m in ("GFA", "SA"):
        j[f"{m}_diff"] = j[f"{m}_last"] - j[f"{m}_prev"]
        s = j[f"{m}_diff"].sort_values()
        out[f"{m}_down"] = [(i, int(v)) for i, v in s.head(top).items() if v < 0]
        out[f"{m}_up"] = [(i, int(v)) for i, v in s.tail(top)[::-1].items() if v > 0]
    prior = adv[adv["날짜"] < last].groupby("adv")[["GFA", "SA"]].sum()
    for m in ("GFA", "SA"):
        was0 = prior[prior[m] > 0].index
        out[f"{m}_new"] = [(i, int(v)) for i, v in j[f"{m}_last"].items()
                           if v > 0 and i not in was0]
    return out


# ────────────────────────── 시트 입출력 ──────────────────────────
def _svc():
    from google.oauth2.credentials import Credentials
    from google.auth.transport.requests import Request
    from googleapiclient.discovery import build
    creds = Credentials.from_authorized_user_file(TOKEN_PATH)  # scopes 생략(공유토큰 규칙)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        with open(TOKEN_PATH, "w", encoding="utf-8") as f:
            f.write(creds.to_json())
    return build("sheets", "v4", credentials=creds)


def _date_in(row, maxc=3):
    for j in range(min(maxc, len(row))):
        v = row[j]
        if isinstance(v, (int, float)) and 44000 < v < 48000:
            return EPOCH + timedelta(days=int(v)), j
    return None, None


def _find_layout(rows):
    """일별표 헤더 행과 GFA/SA/테무 컬럼 인덱스를 동적으로 찾는다.

    대시보드 영역에도 'GFA' 라벨이 있어서, 바로 아래 몇 줄에 날짜 시리얼이
    오는 행만 헤더로 인정한다(월마다 표 위치가 29·30·39행 등으로 다름).
    """
    for i, row in enumerate(rows):
        if not any(isinstance(c, str) and c.strip() == "GFA" for c in row):
            continue
        if not any(_date_in(rows[k])[0] for k in range(i + 1, min(i + 4, len(rows)))):
            continue
        cols = {}
        for j, c in enumerate(row):
            if not isinstance(c, str):
                continue
            s = c.strip().replace(" ", "")
            if s == "GFA":
                cols["GFA"] = j
            elif s.startswith("SA"):
                cols["SA"] = j
            elif s in ("테무_GFA", "테무GFA"):
                cols["테무GFA"] = j
            elif s in ("테무_SA", "테무SA"):
                cols["테무SA"] = j
        return i, cols
    raise RuntimeError("일별표 헤더를 찾지 못했다")


def _tab_for(d: date, titles) -> str:
    pats = [rf"^{d.year % 100}년\s*{d.month}월\s*(네이버\s*)?매출 현황$"]
    for t in titles:
        if any(re.match(p, t.strip()) for p in pats):
            return t
    raise RuntimeError(f"{d.year}-{d.month:02d} 탭을 찾지 못했다: {titles}")


def _a1(col_idx0: int) -> str:
    s, n = "", col_idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def read_sheet_series(year: int, month: int) -> dict:
    """해당 월 탭의 일별 GFA/SA/테무GFA/테무SA를 읽어온다(대조·전월비교용)."""
    svc = _svc()
    titles = [s["properties"]["title"] for s in svc.spreadsheets().get(
        spreadsheetId=SHEET_ID, fields="sheets(properties(title))").execute()["sheets"]]
    tab = _tab_for(date(year, month, 1), titles)
    rows = svc.spreadsheets().values().get(
        spreadsheetId=SHEET_ID, range=f"'{tab}'!A1:V90",
        valueRenderOption="UNFORMATTED_VALUE").execute().get("values", [])
    hidx, cols = _find_layout(rows)
    out = {}
    for row in rows[hidx + 1:]:
        d, _ = _date_in(row)
        if d is None or (d.year, d.month) != (year, month):
            continue
        rec = {}
        for k, j in cols.items():
            v = row[j] if j < len(row) else None
            rec[k] = float(v) if isinstance(v, (int, float)) else None
        out[d.isoformat()] = rec
    return out


def update_sheet(daily: dict, dry_run: bool = True) -> list:
    """일별 GFA/SA/테무GFA/테무SA **4개 컬럼만** 해당 날짜 행에 쓴다.

    다른 칸(목표·수수료·메모 등)은 절대 건드리지 않는다. 날짜별·컬럼별로
    개별 셀 업데이트를 batch로 보내므로 범위 클리어가 일어나지 않는다.
    """
    svc = _svc()
    titles = [s["properties"]["title"] for s in svc.spreadsheets().get(
        spreadsheetId=SHEET_ID, fields="sheets(properties(title))").execute()["sheets"]]

    plan, cache = [], {}
    for ds, vals in sorted(daily.items()):
        d = date.fromisoformat(str(ds)[:10])
        tab = _tab_for(d, titles)
        if tab not in cache:
            rows = svc.spreadsheets().values().get(
                spreadsheetId=SHEET_ID, range=f"'{tab}'!A1:V90",
                valueRenderOption="UNFORMATTED_VALUE").execute().get("values", [])
            hidx, cols = _find_layout(rows)
            rowmap = {}
            for i, row in enumerate(rows[hidx + 1:], start=hidx + 2):
                dd, _ = _date_in(row)
                if dd:
                    rowmap[dd.isoformat()] = (i, row)
            cache[tab] = (cols, rowmap)
        cols, rowmap = cache[tab]
        if d.isoformat() not in rowmap:
            raise RuntimeError(f"{tab}에 {d} 행이 없다")
        rno, cur = rowmap[d.isoformat()]
        for key in ("GFA", "SA", "테무GFA", "테무SA"):
            if key not in cols:
                continue
            j = cols[key]
            old = cur[j] if j < len(cur) else None
            new = int(vals[key])
            plan.append({"tab": tab, "cell": f"{_a1(j)}{rno}", "key": key,
                         "date": d.isoformat(), "old": old, "new": new})

    if not dry_run:
        svc.spreadsheets().values().batchUpdate(
            spreadsheetId=SHEET_ID,
            body={"valueInputOption": "RAW",
                  "data": [{"range": f"'{p['tab']}'!{p['cell']}", "values": [[p["new"]]]}
                           for p in plan]}).execute()
    return plan


# ────────────────────────── 리포트 ──────────────────────────
def build_excel(facts: dict, narrative: dict, out_path: str) -> str:
    """일단위 요약 xlsx 생성. narrative는 섹션별 문단 dict(아래 키 참고)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    daily = facts["daily"]
    dates = sorted(daily)
    HDR = PatternFill("solid", fgColor="4F81BD")
    WB_ = Font(bold=True, color="FFFFFF")
    BOLD = Font(bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "일단위 요약"
    r = 1
    ws.cell(r, 1, f"● 네이버 일단위 매출 요약 ({dates[0]} ~ {dates[-1]})").font = Font(bold=True, size=13)

    def section(title):
        nonlocal r
        r += 2
        ws.cell(r, 1, title).font = BOLD

    def para(text):
        nonlocal r
        r += 1
        ws.cell(r, 1, text)

    def table(headers, rows_):
        nonlocal r
        r += 1
        for j, h in enumerate(headers, start=1):
            c = ws.cell(r, j, h)
            c.fill, c.font, c.alignment = HDR, WB_, Alignment(horizontal="center")
        for row in rows_:
            r += 1
            for j, v in enumerate(row, start=1):
                c = ws.cell(r, j, v)
                if isinstance(v, (int, float)):
                    c.number_format = "#,##0"

    section("[1] 날짜별 상세 (단위: 원)")
    body = []
    for d in dates:
        v = daily[d]
        dt = date.fromisoformat(str(d)[:10])
        body.append([f"{d}({WD[dt.weekday()]})", v["GFA"], v["파워링크"], v["쇼핑검색"],
                     v["SA"], v["테무GFA"], v["테무SA"], v["GFA"] + v["SA"]])
    tot = ["합계"] + [sum(x[i] for x in body) for i in range(1, 8)]
    table(["날짜", "GFA", "파워링크", "쇼핑검색", "SA(파워링크+쇼핑검색)",
           "테무GFA", "테무SA", "합계(GFA+SA)"], body + [tot])
    for j in range(1, 9):
        ws.cell(r, j).font = BOLD

    for key in ("보정", "전일대비", "신규", "페이스", "전망", "테무", "전월비"):
        blk = narrative.get(key)
        if not blk:
            continue
        section(blk["title"])
        if blk.get("headers"):
            table(blk["headers"], blk["rows"])
        for p in blk.get("paras", []):
            para(p)

    for rr in range(1, r + 1):
        ws.row_dimensions[rr].height = 16.5
    for col, w in {"A": 26, "B": 20, "C": 18, "D": 18, "E": 22,
                   "F": 18, "G": 16, "H": 18}.items():
        ws.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def default_out_path(facts: dict) -> str:
    dates = sorted(facts["daily"])
    a, b = dates[0].replace("-", ""), dates[-1].replace("-", "")
    return os.path.join(OUT_DIR, f"네이버_일단위_매출_요약_{a}_{b}.xlsx")


if __name__ == "__main__":
    import sys
    f = compute_from_csv(sys.argv[1])
    print(json.dumps({k: {kk: int(vv) for kk, vv in v.items()} for k, v in f["daily"].items()},
                     ensure_ascii=False, indent=2))
